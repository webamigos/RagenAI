#!/usr/bin/env python3
"""
Check the built corpus: every file yields text, and the documents agree.

    pip install pdfminer.six openpyxl python-docx
    python3 scripts/demo-corpus/verify.py

The second half is the one that matters. A demo fails in front of a prospect
when two documents give two answers to the same question, and that is exactly
the kind of drift an edit to one file introduces silently. Every fact asserted
here is one a demo question in README.md asks about.

Text extraction is also a real check, not a formality: a PDF whose diacritics
are mangled indexes as a different language than the question, and nothing
downstream reports it — the answers just get worse.
"""

import re
import sys
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pdfminer.high_level import extract_text

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from data import (  # noqa: E402
    EXECUTION_2025,
    EXPORT_SHARE,
    MARKETING_BUDGET_2026,
    PRODUCTS,
    REVENUE_2025,
    SLA_FEE_MINIMUM,
)

FILES = HERE / "files"
POLISH_LETTERS = set("ąćęłńóśźżĄĆĘŁŃÓŚŹŻ")

failures: list[str] = []


def check(condition: bool, description: str) -> None:
    if condition:
        print(f"  ok    {description}")
    else:
        print(f"  FAIL  {description}")
        failures.append(description)


def squash(text: str) -> str:
    """Collapse whitespace, so a value broken across PDF lines still matches."""
    return re.sub(r"\s+", " ", text)


def pdf_text(name: str) -> str:
    return squash(extract_text(FILES / name))


def docx_text(name: str) -> str:
    document = Document(FILES / name)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return squash("\n".join(parts))


def xlsx_text(name: str) -> str:
    workbook = load_workbook(FILES / name, data_only=True)
    parts = []
    for sheet in workbook.worksheets:
        parts.append(sheet.title)
        for row in sheet.iter_rows(values_only=True):
            parts.extend(str(value) for value in row if value is not None)
    return squash(" ".join(parts))


def xlsx_column(name: str, sheet_name: str, header: str) -> list:
    """Every value under `header` on `sheet_name`, in order."""
    sheet = load_workbook(FILES / name, data_only=True)[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    header_index = next(
        (index for index, row in enumerate(rows) if header in [str(value) for value in row]),
        None,
    )
    if header_index is None:
        raise AssertionError(f"{name}/{sheet_name}: no column named {header!r}")
    column = [str(value) for value in rows[header_index]].index(header)
    return [row[column] for row in rows[header_index + 1:] if row[column] is not None]


def main() -> int:
    text = {}

    print("Extraction")
    for path in sorted(FILES.iterdir()):
        if path.suffix == ".pdf":
            body = pdf_text(path.name)
        elif path.suffix == ".docx":
            body = docx_text(path.name)
        elif path.suffix == ".xlsx":
            body = xlsx_text(path.name)
        else:
            continue
        text[path.name] = body
        check(len(body) > 1500, f"{path.name} yields readable text ({len(body)} chars)")

    print("\nEncoding")
    for name, body in text.items():
        if name.startswith("employee-handbook"):
            continue
        check(
            bool(POLISH_LETTERS & set(body)),
            f"{name} keeps its Polish diacritics",
        )
    check(
        "gęślą" not in text["employee-handbook-2026.pdf"],
        "employee-handbook-2026.pdf is the English one",
    )

    print("\nProducts agree across the catalogue, the price list and the stock report")
    catalogue = text["katalog-produktow-2026.pdf"]
    prices = text["cennik-2026.xlsx"]
    stock = text["stany-magazynowe-2026-03.xlsx"]
    for sku, name, _unit, _price, _warranty, _lead, _line in PRODUCTS:
        check(
            sku in catalogue and sku in prices and sku in stock,
            f"{sku} ({name}) is in all three",
        )

    print("\nFigures agree")
    total = sum(REVENUE_2025.values())
    monthly = xlsx_column("sprzedaz-2025.xlsx", "Wg miesięcy", "Razem")
    check(
        sum(value for value in monthly if isinstance(value, (int, float))) == 2 * total,
        f"sprzedaz-2025: months plus the total row come to twice {total:,} zł".replace(",", " "),
    )
    check(
        "48 600 000" in text["raport-roczny-2025.pdf"],
        "raport-roczny-2025 quotes the same 48 600 000 zł",
    )
    by_region = xlsx_column("sprzedaz-2025.xlsx", "Wg regionów", "Razem")
    export_row = [value for value in by_region if isinstance(value, (int, float))][4]
    check(
        f"{export_row / total:.0%}" == EXPORT_SHARE,
        f"sprzedaz-2025 puts export at {export_row / total:.0%}, "
        f"which is what raport-roczny-2025 claims ({EXPORT_SHARE})",
    )

    budget_rows = xlsx_column("budzet-marketingowy-2026.xlsx", "Budżet 2026", "Razem")
    numeric = [value for value in budget_rows if isinstance(value, (int, float))]
    check(
        sum(numeric) == 2 * MARKETING_BUDGET_2026,
        "budzet-marketingowy-2026: channels plus the total row come to twice the budget",
    )
    check(
        "1 240 000" in text["protokol-zarzadu-2026-02.docx"],
        "protokol-zarzadu-2026-02 quotes the same 1 240 000 zł (uchwała 1/2026)",
    )

    print("\nRules agree")
    check(
        all(
            warranty == (24 if line == "Regały" else 12)
            for _sku, _name, _unit, _price, warranty, _lead, line in PRODUCTS
        ),
        "every warranty matches § 6 of the framework agreement (24 on racking, 12 otherwise)",
    )
    faq = text["faq-dzial-handlowy.docx"]
    check(
        SLA_FEE_MINIMUM.replace(" ", " ") in squash(text["umowa-serwisowa-sla.pdf"])
        and SLA_FEE_MINIMUM in faq,
        "the SLA minimum fee is the same in the service agreement and the FAQ",
    )
    for document in ("cennik-2026.xlsx", "umowa-ramowa-dostawy.docx", "katalog-produktow-2026.pdf"):
        check("12%" in text[document], f"{document} carries the 50–99 szt. discount (12%)")
    check(
        "14 dni kalendarzowych" in text["procedura-reklamacyjna.docx"]
        and "14 dni kalendarzowych" in faq,
        "the 14-day complaint deadline is the same in the procedure and the FAQ",
    )
    check(
        "1 czerwca 2026" in text["protokol-zarzadu-2026-02.docx"]
        and "1 czerwca 2026" in stock
        and "1 czerwca 2026" in faq,
        "the Gdańsk warehouse opens on the same date in all three documents",
    )
    check(
        "35 dni roboczych" in text["protokol-zarzadu-2026-02.docx"] and "35 dni" in faq,
        "the WZE-20 lead time is the same in the minutes and the FAQ",
    )
    check(
        "poniżej minimum" in stock and "WZE-20" in stock and "AKS-025" in stock,
        "the stock report flags the two SKUs below their minimum level",
    )
    check(
        "ZAM/2026/041" in stock and "10 kwietnia 2026" in stock,
        "the stock report says when the missing WZE-20 units arrive",
    )
    # Read as values, not as text: openpyxl hands back the raw number, while the
    # ingest sees what SheetJS renders through the cell's number format.
    actuals = xlsx_column("budzet-marketingowy-2026.xlsx", "Wykonanie 2025", "Wykonanie 2025")
    channel_actuals = [value for value in actuals if isinstance(value, (int, float))]
    check(
        sum(channel_actuals) == 2 * sum(spent for _, _, spent in EXECUTION_2025),
        "budzet-marketingowy-2026 carries the 2025 actuals, and they add up",
    )

    print("\nThe English handbook agrees with the Polish markdown in scripts/test-env")
    handbook = text["employee-handbook-2026.pdf"]
    markdown = squash(
        "\n".join(
            path.read_text(encoding="utf-8")
            for path in sorted((HERE.parent / "test-env" / "sample-docs").glob("*.md"))
        )
    )
    for fact, description in (
        ("12", "12 remote days a month"),
        ("180", "the PLN 180 remote-work allowance"),
        ("26", "26 days of annual leave"),
        ("450", "the PLN 450 accommodation limit"),
        ("1.15", "the mileage rate"),
    ):
        polish = fact.replace(".", ",")
        check(
            fact in handbook and (fact in markdown or polish in markdown),
            f"{description} is the same in both corpora",
        )

    print()
    if failures:
        print(f"{len(failures)} check(s) failed")
        return 1
    print(f"all {len(text)} files checked, no contradictions")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
