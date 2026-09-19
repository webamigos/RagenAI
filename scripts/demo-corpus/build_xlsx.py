"""
The four spreadsheets.

No formulas anywhere, on purpose: the worker reads XLSX through SheetJS, which
returns a formula cell's *cached* value, and openpyxl writes no cache. A
`=SUM(...)` would therefore reach the index as an empty cell and every question
about a total would be unanswerable. Totals are computed here and written as
numbers.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import build_numbers as n
from data import (
    EXECUTION_2025,
    INBOUND_ORDERS,
    COMPANY,
    DISCOUNTS,
    MARKETING_BUDGET_2026,
    PAYMENT_DAYS,
    PRICE_INCREASE_2026,
    PRODUCTS,
    REGIONS,
    REVENUE_2025,
    VAT_RATE,
    WAREHOUSES,
)

HEADER_FILL = PatternFill("solid", fgColor="252D53")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14, color="252D53")
TOTAL_FONT = Font(bold=True)
THIN = Side(style="thin", color="D4D7E3")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MONEY = '# ##0" zł"'
MONEY_DEC = '# ##0.00" zł"'


def _title(ws, text: str, subtitle: str | None = None) -> int:
    """Write the sheet's title block; return the first free row."""
    ws["A1"] = text
    ws["A1"].font = TITLE_FONT
    row = 2
    if subtitle:
        ws[f"A{row}"] = subtitle
        ws[f"A{row}"].font = Font(size=10, color="5B6178")
        row += 1
    return row + 1


def _header(ws, row: int, labels: list[str]) -> None:
    for column, label in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=column, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BOX
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _widths(ws, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def price_list(path: str) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    for line in ("Regały", "Wózki", "Akcesoria"):
        ws = wb.create_sheet(line)
        row = _title(
            ws,
            f"Cennik 2026 — {line}",
            f"{COMPANY['name']}, obowiązuje od 1 stycznia 2026. Ceny netto w PLN, "
            f"VAT {VAT_RATE}. Podwyżka cen katalogowych: {PRICE_INCREASE_2026}.",
        )
        _header(
            ws,
            row,
            [
                "SKU",
                "Nazwa produktu",
                "Jednostka",
                "Cena netto",
                "Cena brutto",
                "Gwarancja (mies.)",
                "Czas dostawy (dni rob.)",
            ],
        )
        _widths(ws, [12, 42, 11, 14, 14, 17, 20])
        current = row + 1
        for sku, name, unit, price, warranty, lead, product_line in PRODUCTS:
            if product_line != line:
                continue
            ws.cell(row=current, column=1, value=sku)
            ws.cell(row=current, column=2, value=name)
            ws.cell(row=current, column=3, value=unit)
            net = ws.cell(row=current, column=4, value=price)
            net.number_format = MONEY_DEC
            gross = ws.cell(row=current, column=5, value=round(price * 1.23, 2))
            gross.number_format = MONEY_DEC
            ws.cell(row=current, column=6, value=warranty)
            ws.cell(row=current, column=7, value=lead)
            for column in range(1, 8):
                ws.cell(row=current, column=column).border = BOX
            current += 1

    ws = wb.create_sheet("Rabaty i warunki")
    row = _title(ws, "Rabaty progowe i warunki handlowe", COMPANY["name"])
    _header(ws, row, ["Wielkość zamówienia (jedna pozycja)", "Rabat od ceny katalogowej"])
    _widths(ws, [38, 28])
    current = row + 1
    for threshold, discount in DISCOUNTS:
        ws.cell(row=current, column=1, value=threshold).border = BOX
        ws.cell(row=current, column=2, value=discount).border = BOX
        current += 1

    current += 1
    terms = [
        ("Termin płatności", f"{PAYMENT_DAYS} dni od daty wystawienia faktury"),
        ("Waluta rozliczeń", "PLN; dla klientów zagranicznych EUR po kursie NBP z dnia wystawienia faktury"),
        ("Warunki dostawy", "DAP zgodnie z Incoterms 2020, dla zamówień powyżej 15 000 zł netto transport gratis"),
        ("Minimum logistyczne", "2 500 zł netto na zamówienie"),
        ("Ważność oferty", "30 dni od daty wystawienia"),
        ("Montaż", "12% wartości netto zamówionych regałów, wycena indywidualna powyżej 200 szt."),
        ("Przegląd okresowy regałów", "1 490 zł netto za lokalizację, wymagany raz na 12 miesięcy"),
    ]
    ws.cell(row=current, column=1, value="Warunek").fill = HEADER_FILL
    ws.cell(row=current, column=1).font = HEADER_FONT
    ws.cell(row=current, column=2, value="Zasada").fill = HEADER_FILL
    ws.cell(row=current, column=2).font = HEADER_FONT
    current += 1
    for label, value in terms:
        ws.cell(row=current, column=1, value=label).border = BOX
        cell = ws.cell(row=current, column=2, value=value)
        cell.border = BOX
        cell.alignment = Alignment(wrap_text=True)
        current += 1

    wb.save(path)


def sales_report(path: str) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    monthly = n.monthly_by_segment()
    segments = list(REVENUE_2025)

    ws = wb.create_sheet("Podsumowanie")
    row = _title(
        ws,
        "Sprzedaż 2025 — podsumowanie",
        f"{COMPANY['name']}, wartości w PLN netto. Dane skonsolidowane, "
        "zgodne z raportem rocznym 2025.",
    )
    _header(ws, row, ["Segment", "Przychód 2025", "Udział w przychodzie"])
    _widths(ws, [22, 18, 22])
    current = row + 1
    total = sum(REVENUE_2025.values())
    for segment in segments:
        ws.cell(row=current, column=1, value=segment).border = BOX
        value = ws.cell(row=current, column=2, value=REVENUE_2025[segment])
        value.number_format = MONEY
        value.border = BOX
        share = ws.cell(row=current, column=3, value=REVENUE_2025[segment] / total)
        share.number_format = "0.0%"
        share.border = BOX
        current += 1
    ws.cell(row=current, column=1, value="Razem").font = TOTAL_FONT
    grand = ws.cell(row=current, column=2, value=total)
    grand.number_format = MONEY
    grand.font = TOTAL_FONT
    ws.cell(row=current, column=3, value=1).number_format = "0.0%"
    ws.cell(row=current, column=3).font = TOTAL_FONT

    ws = wb.create_sheet("Wg miesięcy")
    row = _title(ws, "Sprzedaż 2025 według miesięcy i segmentów", "Wartości w PLN netto.")
    _header(ws, row, ["Miesiąc", *segments, "Razem"])
    _widths(ws, [16, 15, 15, 15, 15, 16])
    current = row + 1
    for index, month in enumerate(n.MONTHS):
        ws.cell(row=current, column=1, value=month).border = BOX
        month_total = 0
        for offset, segment in enumerate(segments, start=2):
            value = monthly[segment][index]
            month_total += value
            cell = ws.cell(row=current, column=offset, value=value)
            cell.number_format = MONEY
            cell.border = BOX
        cell = ws.cell(row=current, column=len(segments) + 2, value=month_total)
        cell.number_format = MONEY
        cell.font = TOTAL_FONT
        cell.border = BOX
        current += 1
    ws.cell(row=current, column=1, value="Razem 2025").font = TOTAL_FONT
    for offset, segment in enumerate(segments, start=2):
        cell = ws.cell(row=current, column=offset, value=sum(monthly[segment]))
        cell.number_format = MONEY
        cell.font = TOTAL_FONT
    cell = ws.cell(row=current, column=len(segments) + 2, value=total)
    cell.number_format = MONEY
    cell.font = TOTAL_FONT

    ws = wb.create_sheet("Wg regionów")
    row = _title(
        ws,
        "Sprzedaż 2025 według regionów i kwartałów",
        "Wartości w PLN netto. Region Eksport obejmuje Niemcy, Czechy i Litwę.",
    )
    _header(ws, row, ["Region", *n.QUARTERS, "Razem"])
    _widths(ws, [16, 15, 15, 15, 15, 16])
    current = row + 1
    by_region = n.quarterly_by_region()
    quarter_totals = [0, 0, 0, 0]
    for region in REGIONS:
        ws.cell(row=current, column=1, value=region).border = BOX
        for index, value in enumerate(by_region[region]):
            quarter_totals[index] += value
            cell = ws.cell(row=current, column=index + 2, value=value)
            cell.number_format = MONEY
            cell.border = BOX
        cell = ws.cell(row=current, column=6, value=sum(by_region[region]))
        cell.number_format = MONEY
        cell.font = TOTAL_FONT
        cell.border = BOX
        current += 1
    ws.cell(row=current, column=1, value="Razem").font = TOTAL_FONT
    for index, value in enumerate(quarter_totals):
        cell = ws.cell(row=current, column=index + 2, value=value)
        cell.number_format = MONEY
        cell.font = TOTAL_FONT
    cell = ws.cell(row=current, column=6, value=total)
    cell.number_format = MONEY
    cell.font = TOTAL_FONT

    ws = wb.create_sheet("Najwięksi klienci")
    row = _title(
        ws,
        "Najwięksi klienci 2025",
        "Wartości w PLN netto. Żaden klient nie przekracza 10% przychodu — "
        "próg koncentracji przyjęty przez zarząd.",
    )
    _header(ws, row, ["Klient", "Region", "Przychód 2025", "Udział w przychodzie"])
    _widths(ws, [36, 14, 18, 22])
    current = row + 1
    from data import TOP_CUSTOMERS_2025

    for name, region, value in TOP_CUSTOMERS_2025:
        ws.cell(row=current, column=1, value=name).border = BOX
        ws.cell(row=current, column=2, value=region).border = BOX
        cell = ws.cell(row=current, column=3, value=value)
        cell.number_format = MONEY
        cell.border = BOX
        share = ws.cell(row=current, column=4, value=value / total)
        share.number_format = "0.0%"
        share.border = BOX
        current += 1

    wb.save(path)


def stock_report(path: str) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Stany magazynowe")
    row = _title(
        ws,
        "Stany magazynowe — 1 marca 2026",
        f"{COMPANY['name']}. Stan poniżej zapasu minimalnego oznaczony w kolumnie "
        '"Status".',
    )
    _header(
        ws,
        row,
        [
            "SKU",
            "Nazwa produktu",
            "M1 Poznań",
            "M2 Wrocław",
            "M3 Gdańsk",
            "Razem",
            "Zapas min.",
            "Status",
            "Czas dostawy (dni rob.)",
            "Ostatnia inwentaryzacja",
        ],
    )
    _widths(ws, [12, 40, 12, 13, 12, 10, 12, 22, 20, 22])
    current = row + 1
    for sku, name, _unit, _price, _warranty, lead, _line in PRODUCTS:
        stock = n.STOCK[sku]
        total = stock["M1"] + stock["M2"] + stock["M3"]
        status = "poniżej minimum" if total < stock["min"] else "w normie"
        values = [
            sku,
            name,
            stock["M1"],
            stock["M2"],
            stock["M3"],
            total,
            stock["min"],
            status,
            lead,
            stock["inwentaryzacja"],
        ]
        for column, value in enumerate(values, start=1):
            cell = ws.cell(row=current, column=column, value=value)
            cell.border = BOX
            if column == 8 and status != "w normie":
                cell.font = Font(bold=True, color="CB1D3D")
        current += 1

    ws = wb.create_sheet("Zamówienia w drodze")
    row = _title(
        ws,
        "Zamówienia u dostawców — stan na 1 marca 2026",
        "Pozycje zamówione i jeszcze nieprzyjęte. Kolumna „Planowana dostawa” "
        "to termin potwierdzony przez dostawcę.",
    )
    _header(
        ws,
        row,
        ["Nr zamówienia", "SKU", "Nazwa produktu", "Ilość", "Magazyn docelowy",
         "Dostawca", "Planowana dostawa", "Status"],
    )
    _widths(ws, [16, 12, 38, 9, 18, 26, 20, 18])
    current = row + 1
    for values in INBOUND_ORDERS:
        for column, value in enumerate(values, start=1):
            ws.cell(row=current, column=column, value=value).border = BOX
        current += 1

    ws = wb.create_sheet("Magazyny")
    row = _title(ws, "Magazyny", COMPANY["name"])
    _header(ws, row, ["Kod", "Lokalizacja", "Adres", "Status"])
    _widths(ws, [10, 18, 28, 34])
    current = row + 1
    for code, city, address, status in WAREHOUSES:
        for column, value in enumerate((code, city, address, status), start=1):
            ws.cell(row=current, column=column, value=value).border = BOX
        current += 1

    wb.save(path)


def marketing_budget(path: str) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    budget = n.marketing_budget()

    ws = wb.create_sheet("Budżet 2026")
    row = _title(
        ws,
        "Budżet marketingowy 2026",
        "Zatwierdzony uchwałą zarządu nr 1/2026 z 12 lutego 2026. "
        "Wartości w PLN netto.",
    )
    _header(ws, row, ["Kanał", *n.QUARTERS, "Razem", "Udział"])
    _widths(ws, [38, 14, 14, 14, 14, 16, 12])
    current = row + 1
    quarter_totals = [0, 0, 0, 0]
    for channel, values in budget.items():
        ws.cell(row=current, column=1, value=channel).border = BOX
        for index, value in enumerate(values):
            quarter_totals[index] += value
            cell = ws.cell(row=current, column=index + 2, value=value)
            cell.number_format = MONEY
            cell.border = BOX
        channel_total = sum(values)
        cell = ws.cell(row=current, column=6, value=channel_total)
        cell.number_format = MONEY
        cell.font = TOTAL_FONT
        cell.border = BOX
        share = ws.cell(row=current, column=7, value=channel_total / MARKETING_BUDGET_2026)
        share.number_format = "0.0%"
        share.border = BOX
        current += 1
    ws.cell(row=current, column=1, value="Razem").font = TOTAL_FONT
    for index, value in enumerate(quarter_totals):
        cell = ws.cell(row=current, column=index + 2, value=value)
        cell.number_format = MONEY
        cell.font = TOTAL_FONT
    cell = ws.cell(row=current, column=6, value=MARKETING_BUDGET_2026)
    cell.number_format = MONEY
    cell.font = TOTAL_FONT

    ws = wb.create_sheet("Wykonanie 2025")
    row = _title(
        ws,
        "Budżet marketingowy 2025 — plan i wykonanie",
        "Wartości w PLN netto. Baza porównawcza dla budżetu 2026.",
    )
    _header(ws, row, ["Kanał", "Plan 2025", "Wykonanie 2025", "Różnica", "Wykonanie planu"])
    _widths(ws, [38, 16, 18, 14, 18])
    current = row + 1
    plan_total = spent_total = 0
    for channel, plan, spent in EXECUTION_2025:
        plan_total += plan
        spent_total += spent
        ws.cell(row=current, column=1, value=channel).border = BOX
        for column, value in ((2, plan), (3, spent), (4, spent - plan)):
            cell = ws.cell(row=current, column=column, value=value)
            cell.number_format = MONEY
            cell.border = BOX
        share = ws.cell(row=current, column=5, value=spent / plan)
        share.number_format = "0.0%"
        share.border = BOX
        current += 1
    ws.cell(row=current, column=1, value="Razem").font = TOTAL_FONT
    for column, value in ((2, plan_total), (3, spent_total), (4, spent_total - plan_total)):
        cell = ws.cell(row=current, column=column, value=value)
        cell.number_format = MONEY
        cell.font = TOTAL_FONT
    cell = ws.cell(row=current, column=5, value=spent_total / plan_total)
    cell.number_format = "0.0%"
    cell.font = TOTAL_FONT

    ws = wb.create_sheet("Zasady")
    row = _title(ws, "Zasady wydatkowania budżetu", COMPANY["name"])
    _header(ws, row, ["Zasada", "Treść"])
    _widths(ws, [32, 76])
    rules = [
        ("Akceptacja", "Wydatek powyżej 20 000 zł netto wymaga akceptacji dyrektora finansowego."),
        ("Przesunięcia", "Przesunięcie środków między kanałami do 10% wartości kanału nie wymaga zgody zarządu."),
        ("Rezerwa", "Niewykorzystane środki kwartału przechodzą na kwartał następny, ale nie na rok następny."),
        ("Raportowanie", "Dyrektor handlowy raportuje wykonanie budżetu na pierwszym posiedzeniu zarządu po zakończeniu kwartału."),
        ("Targi", "Udział w targach Modernlog i Logimat jest wydatkiem obligatoryjnym i nie podlega przesunięciu."),
    ]
    current = row + 1
    for label, text in rules:
        ws.cell(row=current, column=1, value=label).border = BOX
        cell = ws.cell(row=current, column=2, value=text)
        cell.border = BOX
        cell.alignment = Alignment(wrap_text=True)
        current += 1

    wb.save(path)
